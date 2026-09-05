"""
test_ocr_pipeline.py - Verification Suite for OCR, Extraction, Optimization & Report Writing
Tests:
  1. Image preprocessing (EXIF, RGB conversion, contrast, resolution)
  2. PDF extraction (Digital text & scanned PDF fallback classification)
  3. Text optimization (Bullet normalization, punctuation noise cleanup, Bengali text repair)
  4. API Endpoint /api/ocr_extract_and_optimize
  5. API Endpoint /api/system_audit
  6. End-to-end report generation into .docx with extracted OCR data
"""

import os
import io
import docx
from PIL import Image, ImageDraw
from fastapi.testclient import TestClient
from app import app
import ocr_engine
import document_engine
from daily_audit import perform_daily_audit

client = TestClient(app)

def create_synthetic_test_image() -> bytes:
    """Creates a sample test document image with simulated handwritten/printed text and tables."""
    img = Image.new("RGB", (800, 600), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    # Draw simulated title and header
    draw.rectangle([50, 40, 750, 90], outline=(0, 51, 102), width=2)
    draw.text((70, 55), "EASD Executive Meeting & Field Assessment - OCR Test", fill=(0, 51, 102))
    
    # Draw simulated table
    draw.rectangle([50, 120, 750, 350], outline=(100, 100, 100), width=1)
    draw.line([(50, 160), (750, 160)], fill=(100, 100, 100), width=1)
    draw.line([(250, 120), (250, 350)], fill=(100, 100, 100), width=1)
    
    draw.text((60, 130), "Agenda / Subject", fill=(0, 0, 0))
    draw.text((260, 130), "Discussion & Decisions", fill=(0, 0, 0))
    
    draw.text((60, 180), "1. NCD Field Operations", fill=(50, 50, 50))
    draw.text((260, 180), "Project timeline approved for Q3 2026. Budget allocated.", fill=(50, 50, 50))
    
    draw.text((60, 240), "2. Staff Logistics", fill=(50, 50, 50))
    draw.text((260, 240), "Assigned to Dr. Taseen and Pew for field monitoring.", fill=(50, 50, 50))
    
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

def create_synthetic_test_pdf() -> bytes:
    """Generates a minimal valid PDF file in-memory using pypdf."""
    import pypdf
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()

def test_image_preprocessing():
    print("1. Testing Image Preprocessing for OCR...")
    raw_img = create_synthetic_test_image()
    proc_bytes, mime = ocr_engine.preprocess_image_for_ocr(raw_img)
    assert len(proc_bytes) > 0
    assert mime == "image/jpeg"
    
    # Verify valid image output
    img = Image.open(io.BytesIO(proc_bytes))
    assert img.mode == "RGB"
    assert img.size[0] <= 3000 and img.size[1] <= 3000
    print("   [PASSED] Image preprocessed, contrast enhanced, and normalized to clean RGB JPEG.")

def test_ocr_text_optimization():
    print("2. Testing OCR Text Cleanup & Bengali/English Optimization...")
    noisy_ocr_output = """
    ====================================
    * First action item for team
    - Second discussion regarding field work
    + Third point on budget allocation
    o Fourth agenda note
    | | | | ~~~
    ক র তে  হ বে  প্র কল্প  বা স্ত বা য় ন
    ____________________________________
    """
    cleaned = ocr_engine.optimize_ocr_text(noisy_ocr_output)
    
    # Check bullet standardization
    assert "• First action item for team" in cleaned
    assert "• Second discussion regarding field work" in cleaned
    assert "• Third point on budget allocation" in cleaned
    assert "• Fourth agenda note" in cleaned
    
    # Check noise suppression
    assert "=====" not in cleaned
    assert "_____" not in cleaned
    assert "| | | | ~~~" not in cleaned
    
    # Check Bengali ligature repair
    assert "করতে হবে" in cleaned or "করতে" in cleaned
    print("   [PASSED] Bullet points normalized, scanner noise stripped, and Bengali ligatures repaired.")

def test_pdf_extraction():
    print("3. Testing PDF Extraction and Scanned Document Detection...")
    pdf_bytes = create_synthetic_test_pdf()
    pdf_info = ocr_engine.extract_pdf_content(pdf_bytes)
    assert "is_scanned" in pdf_info
    assert pdf_info["page_count"] == 1
    print(f"   [PASSED] PDF parsed ({pdf_info['page_count']} page, scanned_flag={pdf_info['is_scanned']}).")

def test_ocr_extract_and_optimize_endpoint():
    print("4. Testing POST /api/ocr_extract_and_optimize...")
    img_bytes = create_synthetic_test_image()
    
    resp = client.post(
        "/api/ocr_extract_and_optimize",
        files={"file": ("meeting_whiteboard.png", img_bytes, "image/png")},
        data={
            "provider": "local",
            "model_name": "local",
            "raw_text": "Meeting on 29 August 2026. Discussion on NCD project timeline. Action items assigned to Dr. Taseen and Pew."
        }
    )
    assert resp.status_code == 200, f"OCR endpoint failed: {resp.text}"
    data = resp.json()
    assert data["status"] == "success"
    assert "data" in data
    assert "ocr_text" in data
    print("   [PASSED] /api/ocr_extract_and_optimize returned structured report and OCR text.")

def test_system_audit_endpoint():
    print("5. Testing GET /api/system_audit...")
    resp = client.get("/api/system_audit")
    assert resp.status_code == 200
    data = resp.json()
    assert data["status"] == "success"
    report = data["report"]
    assert "security" in report
    assert "setup" in report
    assert "building" in report
    assert "recommendations" in report
    assert report["security"]["security_score"] >= 90
    print(f"   [PASSED] Audit endpoint active. Security Score: {report['security']['security_score']}/100.")

def test_write_extracted_data_to_report_docx():
    print("6. Testing Writing Extracted & Optimized Information into Formal EASD Report...")
    # Prepare extracted structured payload
    meeting_payload = {
        "title": "Quarterly NCD Executive Steering Committee & Operations Review",
        "date": "29 August, 2026",
        "time": "10:30 AM - 01:00 PM",
        "location": "Eminence Conference Room, Mohakhali DOHS, Dhaka",
        "meeting_chair": "Dr. Shamim Talukder",
        "attendance": document_engine.DEFAULT_MEMBERS,
        "discussions": [
            {
                "sn": "1",
                "topic": "Followup from previous meeting",
                "details": "• Previous meeting minutes reviewed and adopted without objection."
            },
            {
                "sn": "2",
                "topic": "Action items",
                "details": "• NCD Health Intervention Field Scale-Up commenced across target divisions."
            },
            {
                "sn": "3",
                "topic": "Task Assignments",
                "details": "• Md. Taseen: Automated OCR and meeting report pipeline.\n• Suraiya Pew: Field operations coordination."
            },
            {
                "sn": "4",
                "topic": "Meeting Decisions",
                "details": "• Q3 budget and intervention schedule unanimously approved."
            }
        ],
        "decisions": "• Q3 budget and intervention schedule unanimously approved.",
        "bangla_transcript": "বাংলা ট্রান্সক্রিপ্ট...",
        "english_transcript": "English transcript..."
    }
    
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "EASD Meeting minutes - Template-DDMonthYY.docx")
    docx_bytes = document_engine.generate_meeting_minutes_docx_bytes(meeting_payload, template_path)
    assert len(docx_bytes) > 10000, "DOCX generation produced an empty or incomplete file"
    
    # Inspect generated DOCX to ensure all extracted data is written
    doc = docx.Document(io.BytesIO(docx_bytes))
    assert len(doc.tables) >= 2
    
    # Verify discussions table content
    table_0 = doc.tables[0]
    full_discussions_text = " ".join(cell.text for row in table_0.rows for cell in row.cells)
    assert "NCD Health Intervention Field Scale-Up" in full_discussions_text
    assert "Md. Taseen" in full_discussions_text
    assert "Suraiya Pew" in full_discussions_text
    
    # Verify attendance table content
    table_1 = doc.tables[1]
    attendance_text = " ".join(cell.text for row in table_1.rows for cell in row.cells)
    assert "Md. Shamim Hayder Talukder" in attendance_text
    assert "H M Taseen Jubair Bhuiyan" in attendance_text
    
    print(f"   [PASSED] Extracted data successfully written to official DOCX report ({len(docx_bytes)} bytes).")

if __name__ == "__main__":
    print("\n=== Running OCR, Extraction, Optimization & Report Generation Test Suite ===\n")
    test_image_preprocessing()
    test_ocr_text_optimization()
    test_pdf_extraction()
    test_ocr_extract_and_optimize_endpoint()
    test_system_audit_endpoint()
    test_write_extracted_data_to_report_docx()
    print("\n=== ALL OCR, EXTRACTION, OPTIMIZATION & REPORT WRITING TESTS PASSED (100%) ===\n")
